"""Check that a source document really is the paper we think it is.

Matching on subject, school and year alone trusts a folder name. Before
writing answers into a paper, look for the paper's own question text inside
the document: if the corpus says question 5 asks about atomic absorption
spectrometers, those words should be somewhere in the file the answers came
from. Papers whose questions are pure images have no text to match, so a low
score is "cannot tell", not "wrong" — the caller decides."""
import re, json, os, sys

def norm(s):
    return re.sub(r'[^a-z0-9]+', ' ', (s or '').lower()).strip()

def doc_text(path):
    if path.lower().endswith('.docx'):
        import docx
        d = docx.Document(path)
        bits = [p.text for p in d.paragraphs]
        for t in d.tables:
            for r in t.rows: bits += [c.text for c in r.cells]
        return norm(" ".join(bits))
    import pymupdf
    return norm(" ".join(p.get_text() for p in pymupdf.open(path)))

def shingles(text, n=6):
    w = text.split()
    return {" ".join(w[i:i+n]) for i in range(0, max(0, len(w)-n+1))}

def score(question_texts, path):
    """How many of the paper's questions can be found in the document."""
    hay = doc_text(path)
    hs = shingles(hay)
    found = checked = 0
    for qt in question_texts:
        t = norm(qt)
        t = re.sub(r'^\d{1,2}\s+', '', t)          # drop the leading question number
        w = t.split()
        if len(w) < 8: continue                     # nothing distinctive to match on
        checked += 1
        probes = [" ".join(w[i:i+6]) for i in range(0, min(len(w)-5, 40), 5)]
        if any(p in hs for p in probes): found += 1
    return found, checked
